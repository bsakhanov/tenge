"""Сборка docx из full.md: стили, таблицы, колонтитул, оглавление (двухпроходное).
Запуск: python docx_build.py <1|2> <каталог сборки>. Обычно вызывается из make.py."""
import re, subprocess, sys, json, copy, os
from docx import Document
from docx.shared import Pt, RGBColor, Emu, Mm, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

PASS=int(sys.argv[1]) if len(sys.argv)>1 else 1
os.chdir(sys.argv[2] if len(sys.argv)>2 else ".")
TEXT_W_DXA=9412  # 11906 - 2*1247

def make_ref():
    subprocess.run("pandoc -o ref_raw.docx --print-default-data-file reference.docx",shell=True,check=True)
    d=Document("ref_raw.docx")
    st=d.styles
    class _S:
        def __init__(self,s): self.s=s
        def __getitem__(self,k):
            for x in self.s:
                if x.name==k: return x
            raise KeyError(k)
        def __iter__(self): return iter(self.s)
    st=_S(st)
    def setfont(style,name,size=None,bold=None,color=None,italic=None):
        s=st[style]; f=s.font; f.name=name
        rpr=s.element.get_or_add_rPr()
        rf=rpr.find(qn('w:rFonts'))
        if rf is None:
            rf=OxmlElement('w:rFonts'); rpr.append(rf)
        for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'): rf.set(qn(a),name)
        if size: f.size=Pt(size)
        if bold is not None: f.bold=bold
        if italic is not None: f.italic=italic
        if color: f.color.rgb=RGBColor.from_string(color)
    for s in ("Normal","Body Text","First Paragraph","Compact","Block Text"):
        if s in [x.name for x in st]:
            setfont(s,"Times New Roman",12 if s!="Compact" else 10)
    for s in ("Body Text","First Paragraph"):
        pf=st[s].paragraph_format
        pf.space_before=Pt(0); pf.space_after=Pt(6); pf.line_spacing=1.15
        pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY; pf.widow_control=True
    setfont("Heading 1","Arial",18,True,"1F3864")
    setfont("Heading 2","Arial",14,True,"1F3864")
    setfont("Heading 3","Arial",12,True,"2F5496")
    for s,sb,sa in (("Heading 1",0,12),("Heading 2",14,6),("Heading 3",10,4)):
        pf=st[s].paragraph_format; pf.space_before=Pt(sb); pf.space_after=Pt(sa); pf.keep_with_next=True; pf.keep_together=True
    setfont("Title","Arial",28,True,"1F3864")
    setfont("Subtitle","Arial",15,False,"2F5496")
    st["Title"].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
    st["Subtitle"].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
    st["Title"].paragraph_format.space_before=Pt(int(os.environ.get("TITLE_SPACE_BEFORE","120"))); st["Title"].paragraph_format.space_after=Pt(12)
    for s in ("Footnote Text","Table Caption","Image Caption"):
        if s in [x.name for x in st]: setfont(s,"Times New Roman",10)
    sec=d.sections[0]
    sec.page_width=Emu(int(11906*635)); sec.page_height=Emu(int(16838*635))
    for a in ("left_margin","right_margin"): setattr(sec,a,Emu(1247*635))
    sec.top_margin=Emu(1247*635); sec.bottom_margin=Emu(1247*635)
    d.save("reference.docx")

def prep_md():
    t=open("full.md").read()
    t=t.replace("\\newpage\n","")
    open("full_clean.md","w").write(t)

def pandoc():
    subprocess.run("pandoc full_clean.md -f markdown+pipe_tables+fenced_divs -t docx --reference-doc=reference.docx -o base.docx",shell=True,check=True)

def set_cell_borders(tbl):
    tblPr=tbl._tbl.tblPr
    b=OxmlElement('w:tblBorders')
    for e in ('top','left','bottom','right','insideH','insideV'):
        x=OxmlElement(f'w:{e}'); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:space'),'0'); x.set(qn('w:color'),'8EA0C0'); b.append(x)
    old=tblPr.find(qn('w:tblBorders'))
    if old is not None: tblPr.remove(old)
    tblPr.append(b)
    m=OxmlElement('w:tblCellMar')
    for e,v in (('top',50),('left',90),('bottom',50),('right',90)):
        x=OxmlElement(f'w:{e}'); x.set(qn('w:w'),str(v)); x.set(qn('w:type'),'dxa'); m.append(x)
    old=tblPr.find(qn('w:tblCellMar'))
    if old is not None: tblPr.remove(old)
    tblPr.append(m)

def shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr()
    for o in tcPr.findall(qn('w:shd')): tcPr.remove(o)
    s=OxmlElement('w:shd'); s.set(qn('w:val'),'clear'); s.set(qn('w:color'),'auto'); s.set(qn('w:fill'),fill); tcPr.append(s)

def fmt_tables(d):
    for tbl in d.tables:
        set_cell_borders(tbl)
        tbl.alignment=WD_TABLE_ALIGNMENT.CENTER
        nrows=len(tbl.rows); ncols=len(tbl.columns)
        # widths by text length
        lens=[0]*ncols
        for r in tbl.rows:
            for i,c in enumerate(r.cells[:ncols]):
                L=len(c.text); lens[i]=max(lens[i],min(L,220))
        wts=[max(l,8)**0.7 for l in lens]
        lw=[0]*ncols
        for r in tbl.rows:
            for i,c in enumerate(r.cells[:ncols]):
                for tok in re.split(r"[\s/]+",c.text):
                    lw[i]=max(lw[i],len(tok))
        minw=[min(2300,lw[i]*118+230) for i in range(ncols)]
        tot=sum(wts); widths=[int(TEXT_W_DXA*w/tot) for w in wts]
        for _ in range(5):
            short=[i for i in range(ncols) if widths[i]<minw[i]]
            if not short: break
            deficit=sum(minw[i]-widths[i] for i in short)
            for i in short: widths[i]=minw[i]
            rest=[i for i in range(ncols) if i not in short]
            rsum=sum(widths[i] for i in rest)
            for i in rest: widths[i]=int(widths[i]-deficit*widths[i]/rsum)
        widths[-1]+=TEXT_W_DXA-sum(widths)
        tblPr=tbl._tbl.tblPr
        for o in tblPr.findall(qn('w:tblW')): tblPr.remove(o)
        tw=OxmlElement('w:tblW'); tw.set(qn('w:w'),str(TEXT_W_DXA)); tw.set(qn('w:type'),'dxa'); tblPr.append(tw)
        lay=OxmlElement('w:tblLayout'); lay.set(qn('w:type'),'fixed'); tblPr.append(lay)
        grid=tbl._tbl.tblGrid
        for g,wd in zip(grid.findall(qn('w:gridCol')),widths): g.set(qn('w:w'),str(wd))
        tot_chars=sum(len(c.text) for r in tbl.rows for c in r.cells)
        keep = nrows<=16 and tot_chars<2300
        for ri,r in enumerate(tbl.rows):
            trPr=r._tr.get_or_add_trPr()
            cs=OxmlElement('w:cantSplit'); trPr.append(cs)
            if ri==0:
                th=OxmlElement('w:tblHeader'); trPr.append(th)
            for ci,c in enumerate(r.cells[:ncols]):
                c.width=Emu(widths[ci]*635)
                for p in c.paragraphs:
                    p.paragraph_format.space_after=Pt(0); p.paragraph_format.space_before=Pt(0)
                    p.paragraph_format.line_spacing=1.0
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT
                    if keep and ri<nrows-1: p.paragraph_format.keep_with_next=True
                    for run in p.runs:
                        run.font.size=Pt(9); run.font.name="Times New Roman"
                        if ri==0: run.font.bold=True
                if ri==0: shade(c,'D9E2F3')

def para_keep(d):
    ps=list(d.paragraphs)
    for i in range(len(ps)-1):
        nx=ps[i+1].text.strip()
        if nx.startswith("Таблица ") and len(nx)<160 and ps[i+1].runs and ps[i+1].runs[0].bold and len(ps[i].text)<260:
            ps[i].paragraph_format.keep_with_next=True
    prev=None
    for p in d.paragraphs:
        t=p.text.strip()
        if t.startswith("Таблица ") and len(t)<160 and p.runs and p.runs[0].bold:
            p.paragraph_format.keep_with_next=True
            p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
        if p.style.name=="Heading 1":
            p.paragraph_format.page_break_before=True

def footer(d):
    sec=d.sections[0]
    sec.footer.is_linked_to_previous=False
    p=sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
    for r in list(p.runs): r._r.getparent().remove(r._r)
    p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    run=p.add_run()
    for typ,txt in (("begin",None),("instr"," PAGE "),("separate",None),("text","1"),("end",None)):
        if typ=="instr":
            e=OxmlElement('w:instrText'); e.set(qn('xml:space'),'preserve'); e.text=txt
        elif typ=="text":
            e=OxmlElement('w:t'); e.text=txt
        else:
            e=OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'),typ)
        run._r.append(e)
    run.font.name="Times New Roman"; run.font.size=Pt(10)

def headings(d):
    out=[]
    for p in d.paragraphs:
        if p.style.name in ("Heading 1","Heading 2"):
            out.append((1 if p.style.name=="Heading 1" else 2,p.text.strip()))
    return out

def insert_toc(d,entries,pages):
    # find Heading 1 "Резюме для руководства"
    target=[p for p in d.paragraphs if p.style.name=="Heading 1"][0]
    def new_par_before(text,style=None):
        np_=OxmlElement('w:p')
        target._p.addprevious(np_)
        from docx.text.paragraph import Paragraph
        pp=Paragraph(np_,target._parent)
        return pp
    # TOC title: use Heading 1 look but exclude from outline -> set style Normal with formatting
    h=new_par_before("")
    h.style=d.styles["Normal"]
    pf=h.paragraph_format; pf.page_break_before=True; pf.space_after=Pt(12)
    r=h.add_run("Оглавление"); r.font.name="Arial"; r.font.size=Pt(18); r.font.bold=True; r.font.color.rgb=RGBColor.from_string("1F3864")
    for (lvl,txt),pg in zip(entries,pages):
        p=new_par_before("")
        p.style=d.styles["Normal"]
        pf=p.paragraph_format
        pf.left_indent=Cm(0 if lvl==1 else 0.6); pf.space_after=Pt(2 if lvl==2 else 3); pf.space_before=Pt(6 if lvl==1 else 0)
        pf.alignment=WD_ALIGN_PARAGRAPH.LEFT
        pf.tab_stops.add_tab_stop(Emu(TEXT_W_DXA*635),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
        run=p.add_run(txt+"\t"+str(pg))
        run.font.name="Arial" if lvl==1 else "Times New Roman"; run.font.size=Pt(11 if lvl==1 else 10.5); run.font.bold=(lvl==1)
    # make Резюме start on new page already via Heading1 page_break_before

def props(d):
    cp=d.core_properties
    a=os.environ.get("DOC_AUTHOR","Саханов Бейбит")
    cp.author=a; cp.last_modified_by=a
    cp.title=os.environ.get("DOC_TITLE","Цифровой тенге в Казахстане: доклад-исследование"); cp.subject=os.environ.get("DOC_SUBJECT","Цифровая валюта центрального банка, РК"); cp.language="ru-RU"

def finish(inp,out,pages=None):
    d=Document(inp)
    fmt_tables(d); para_keep(d); footer(d); props(d)
    ents=headings(d)
    # exclude front headings in TOC? keep all H1/H2
    if pages is None: pages=["00"]*len(ents)
    insert_toc(d,ents,pages)
    d.save(out)
    return ents

if __name__=="__main__":
    if PASS==1:
        make_ref(); prep_md(); pandoc()
        ents=finish("base.docx","pass1.docx")
        json.dump(ents,open("toc_entries.json","w"),ensure_ascii=False)
        print("entries",len(ents))
    else:
        ents=json.load(open("toc_entries.json"))
        pages=json.load(open("toc_pages.json"))
        finish("base.docx","final.docx",pages)
        print("final built")
